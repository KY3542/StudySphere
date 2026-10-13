import io
import json
import os
import re
import time
import urllib.parse
import urllib.request
from datetime import datetime

from docx import Document
from google import genai
from google.genai import types
from gtts import gTTS
import graphviz
from pypdf import PdfReader
from pptx import Presentation
from reportlab.lib.colors import HexColor
# --- REPORTLAB PDF GENERATION IMPORTS ---
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
import streamlit as st

# --- PAGE CONFIGURATION ---
st.set_page_config(
    page_title="StudySphere", page_icon="logo.png", layout="wide"
)


# --- SECURE API KEY RETRIEVAL (SERVER-SIDE ONLY) ---
def get_server_api_key():
  """Retrieve Gemini API key from environment or secrets without exposing inputs in UI."""
  if "GEMINI_API_KEY" in os.environ:
    return os.environ["GEMINI_API_KEY"]
  try:
    if "GEMINI_API_KEY" in st.secrets:
      return st.secrets["GEMINI_API_KEY"]
  except Exception:
    pass
  return ""


API_KEY = get_server_api_key()

# --- VIVARILY NEON GLOW THEME & CRISP IMAGE RENDER CSS ---
st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Space+Mono:wght@400;700&family=Inter:wght@400;500;600&display=swap');

    .stApp {
        background-color: #0a050d !important;
        background-image: 
            radial-gradient(circle at 85% 30%, rgba(147, 51, 234, 0.25) 0%, rgba(236, 72, 153, 0.15) 35%, transparent 70%),
            linear-gradient(180deg, #0e0714 0%, #0a050d 100%) !important;
        background-attachment: fixed !important;
        color: #f3e8ff !important;
        font-family: 'Inter', sans-serif !important;
    }

    img {
        image-rendering: -webkit-optimize-contrast !important;
        image-rendering: crisp-edges !important;
        object-fit: contain !important;
    }

    div[data-testid="stSidebarHeader"] {
        display: none !important;
    }

    .circle-logo-container {
        display: flex !important;
        justify-content: center !important;
        align-items: center !important;
        width: 100% !important;
        margin: 10px auto 15px auto !important;
    }

    .circle-logo-container img {
        border-radius: 50% !important;
        object-fit: cover !important;
        border: 2px solid rgba(217, 70, 239, 0.6) !important;
        box-shadow: 0px 0px 20px rgba(217, 70, 239, 0.5) !important;
    }

    h1 {
        font-family: 'Space Mono', monospace !important;
        color: #ffffff !important;
        font-size: 2.2rem !important;
        font-weight: 700 !important;
        letter-spacing: 2px !important;
        text-shadow: 0 0 12px rgba(217, 70, 239, 0.3) !important;
    }

    h2, h3, h4 {
        font-family: 'Space Mono', monospace !important;
        color: #c084fc !important;
        font-weight: 700 !important;
    }

    section[data-testid="stSidebar"] {
        background-color: #0f0817 !important;
        border-right: 1px solid #2e1065 !important;
    }

    p, span, label, div, caption, .stCaption, li {
        color: #e9d5ff !important;
    }

    .stButton>button {
        background: rgba(15, 8, 23, 0.7) !important;
        color: #c084fc !important;
        font-family: 'Space Mono', monospace !important;
        font-size: 0.85rem !important;
        font-weight: 700 !important;
        text-transform: uppercase !important;
        border: 1px solid #d946ef !important;
        border-radius: 4px !important;
        box-shadow: 0px 0px 10px rgba(217, 70, 239, 0.2) !important;
        transition: all 0.2s ease-in-out !important;
    }

    .stButton>button:hover {
        background: #d946ef !important;
        color: #0a050d !important;
        box-shadow: 0px 0px 18px rgba(217, 70, 239, 0.6) !important;
    }

    div[data-baseweb="select"] > div {
        background-color: #140b21 !important;
        border: 1px solid #a855f7 !important;
        color: #ffffff !important;
        font-family: 'Space Mono', monospace !important;
    }

    div[data-testid="stFileUploaderDropzone"] {
        background-color: #140b21 !important;
        border: 1px dashed #d946ef !important;
    }

    div[data-element-container] > div[data-testid="stVerticalBlock"] > div[style*="border"] {
        background: rgba(20, 11, 33, 0.6) !important;
        border: 1px solid #3b0764 !important;
        border-radius: 8px !important;
        backdrop-filter: blur(8px) !important;
    }

    div[data-testid="stExpander"] {
        background-color: #140b21 !important;
        border: 1px solid #3b0764 !important;
    }

    input, textarea {
        background-color: #0f0817 !important;
        border: 1px solid #581c87 !important;
        color: #ffffff !important;
        font-family: 'Space Mono', monospace !important;
    }

    div[data-testid="stChatInput"] textarea {
        background-color: #140b21 !important;
        border: 1px solid #d946ef !important;
        color: #ffffff !important;
    }

    div[data-testid="stChatMessage"] {
        background-color: #140b21 !important;
        border: 1px solid #3b0764 !important;
        border-radius: 6px !important;
    }

    div[data-testid="stChatMessage"]:nth-child(even) {
        background-color: #1e0a2f !important;
        border-color: #7e22ce !important;
    }

    header[data-testid="stHeader"] {
        background: transparent !important;
    }

    div[data-testid="stSegmentedControl"] {
        background-color: #0f0817 !important;
        border: 1px solid #3b0764 !important;
        border-radius: 8px !important;
        padding: 4px !important;
    }

    div[data-testid="stSegmentedControl"] button {
        background: transparent !important;
        color: #c084fc !important;
        font-family: 'Space Mono', monospace !important;
        font-size: 0.8rem !important;
        font-weight: 700 !important;
        border: none !important;
        border-radius: 6px !important;
        transition: all 0.2s ease-in-out !important;
    }

    div[data-testid="stSegmentedControl"] button[aria-selected="true"] {
        background: #d946ef !important;
        color: #0a050d !important;
        box-shadow: 0px 0px 15px rgba(217, 70, 239, 0.6) !important;
    }

    div[data-testid="stSegmentedControl"] button:hover:not([aria-selected="true"]) {
        background: rgba(217, 70, 239, 0.15) !important;
        color: #f472b6 !important;
    }
    </style>
""",
    unsafe_allow_html=True,
)

# --- FILE PERSISTENCE & LOCAL HELPERS ---
SETTINGS_FILE = "settings.txt"
PHOTO_FILE = "profile_pic.png"


def load_local_settings():
  settings = {"user": "Active Student", "questions_answered": 0, "study_time_min": 0}
  if os.path.exists(SETTINGS_FILE):
    try:
      with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
        for line in f:
          if "=" in line:
            k, v = line.strip().split("=", 1)
            if k in ["questions_answered", "study_time_min"]:
              settings[k] = int(v)
            else:
              settings[k] = v
    except Exception:
      pass
  return settings


def save_local_settings(username, q_count, s_time):
  try:
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
      f.write(f"user={username}\n")
      f.write(f"questions_answered={q_count}\n")
      f.write(f"study_time_min={s_time}\n")
  except Exception:
    pass


def clean_graphviz_dot(raw_dot_text: str) -> str:
  """Cleans raw model response and converts single quotes in Graphviz attributes to valid double quotes."""
  if not raw_dot_text:
    return ""

  cleaned = re.sub(r'```(?:dot|graphviz)?', '', raw_dot_text, flags=re.IGNORECASE)
  cleaned = cleaned.replace('```', '').strip()

  start_match = re.search(r'((?:digraph|graph)\s+[\s\S]*)', cleaned, re.IGNORECASE)
  if start_match:
    cleaned = start_match.group(1)

  # CRITICAL FIX: Replace invalid single quotes in DOT attributes with double quotes (e.g. color='#123' -> color="#123")
  cleaned = re.sub(r"=\s*'([^']*)'", r'="\1"', cleaned)

  return cleaned.strip()


def parse_uploaded_document(uploaded_file):
  """Extract text from PDF, DOCX, PPTX, or TXT notes."""
  text = ""
  filename = uploaded_file.name.lower()
  try:
    if filename.endswith(".pdf"):
      reader = PdfReader(io.BytesIO(uploaded_file.read()))
      for page in reader.pages:
        t = page.extract_text()
        if t:
          text += t + "\n"
    elif filename.endswith(".docx"):
      doc = Document(io.BytesIO(uploaded_file.read()))
      for p in doc.paragraphs:
        text += p.text + "\n"
    elif filename.endswith(".pptx"):
      prs = Presentation(io.BytesIO(uploaded_file.read()))
      for slide in prs.slides:
        for shape in slide.shapes:
          if hasattr(shape, "text"):
            text += shape.text + "\n"
    elif filename.endswith(".txt"):
      text = uploaded_file.read().decode("utf-8")
  except Exception as e:
    st.error(f"Error reading file {filename}: {e}")
  return text


def handle_api_error(e):
  """Format user-friendly error message for API limit issues."""
  err_msg = str(e)
  if "429" in err_msg or "RESOURCE_EXHAUSTED" in err_msg:
    return "⏳ **Gemini API Limit Reached**: Daily quota exhausted for this API key. Please try again later or add billing in Google AI Studio."
  return f"API Error: {err_msg}"


# --- REPORT CARD & CHAT LOG PDF GENERATION FUNCTIONS ---
def get_grade_letter(percentage):
  if percentage >= 90:
    return "A+"
  elif percentage >= 80:
    return "A"
  elif percentage >= 70:
    return "B"
  elif percentage >= 60:
    return "C"
  elif percentage >= 50:
    return "D"
  else:
    return "F"


def generate_pdf_report_card(
    username,
    subject_title,
    questions_scores,
    total_score,
    max_score,
    percentage,
    grade,
    real_time_str,
):
  buffer = io.BytesIO()
  doc = SimpleDocTemplate(
      buffer,
      pagesize=letter,
      rightMargin=36,
      leftMargin=36,
      topMargin=36,
      bottomMargin=36,
  )
  story = []
  styles = getSampleStyleSheet()

  title_style = ParagraphStyle(
      "TitleStyle",
      parent=styles["Heading1"],
      fontSize=24,
      leading=28,
      textColor=HexColor("#1e293b"),
      spaceAfter=4,
  )
  sub_title_style = ParagraphStyle(
      "SubTitleStyle",
      parent=styles["Normal"],
      fontSize=12,
      leading=16,
      textColor=HexColor("#64748b"),
      spaceAfter=15,
  )
  meta_style = ParagraphStyle(
      "MetaStyle",
      parent=styles["Normal"],
      fontSize=11,
      leading=16,
      textColor=HexColor("#0f172a"),
  )
  th_style = ParagraphStyle(
      "THStyle",
      parent=styles["Normal"],
      fontSize=10,
      leading=12,
      textColor=HexColor("#ffffff"),
  )
  td_style = ParagraphStyle(
      "TDStyle",
      parent=styles["Normal"],
      fontSize=10,
      leading=14,
      textColor=HexColor("#334155"),
  )

  story.append(Paragraph("<b>REPORT CARD</b>", title_style))
  story.append(
      Paragraph(f"<b>SUBJECT:</b> {subject_title.upper()}", sub_title_style)
  )

  meta_data = [[
      Paragraph(f"<b>Name:</b> {username}", meta_style),
      Paragraph(f"<b>Date:</b> {real_time_str}", meta_style),
  ]]
  meta_table = Table(meta_data, colWidths=[270, 270])
  meta_table.setStyle(
      TableStyle([
          ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
          ("TOPPADDING", (0, 0), (-1, -1), 0),
      ])
  )
  story.append(meta_table)
  story.append(Spacer(1, 10))

  table_data = [[
      Paragraph("<b>QUESTION</b>", th_style),
      Paragraph("<b>MARKS OBTAINED</b>", th_style),
  ]]
  for item in questions_scores:
    table_data.append([
        Paragraph(
            f"Question {item['q_num']}: {item['question_title']}", td_style
        ),
        Paragraph(f"{item['score']} / {item['max_score']}", td_style),
    ])

  q_table = Table(table_data, colWidths=[380, 160])
  q_table.setStyle(
      TableStyle([
          ("BACKGROUND", (0, 0), (-1, 0), HexColor("#1e293b")),
          ("ALIGN", (0, 0), (-1, -1), "LEFT"),
          ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
          ("TOPPADDING", (0, 0), (-1, -1), 8),
          (
              "ROWBACKGROUNDS",
              (0, 1),
              (-1, -1),
              [HexColor("#f8fafc"), HexColor("#ffffff")],
          ),
          ("GRID", (0, 0), (-1, -1), 0.5, HexColor("#e2e8f0")),
      ])
  )
  story.append(q_table)
  story.append(Spacer(1, 15))

  summary_data = [[
      Paragraph(
          f"<b>TOTAL SCORE:</b> {total_score} / {max_score} ({percentage:.1f}%)",
          meta_style,
      ),
      Paragraph(f"<b>FINAL GRADE:</b> {grade}", meta_style),
  ]]
  summary_table = Table(summary_data, colWidths=[300, 240])
  summary_table.setStyle(
      TableStyle([
          ("BACKGROUND", (0, 0), (-1, -1), HexColor("#f1f5f9")),
          ("BOX", (0, 0), (-1, -1), 1, HexColor("#cbd5e1")),
          ("PADDING", (0, 0), (-1, -1), 10),
      ])
  )
  story.append(summary_table)

  doc.build(story)
  buffer.seek(0)
  return buffer.getvalue()


def generate_chat_pdf(username, subject_title, chat_history):
  buffer = io.BytesIO()
  doc = SimpleDocTemplate(
      buffer,
      pagesize=letter,
      rightMargin=36,
      leftMargin=36,
      topMargin=36,
      bottomMargin=36,
  )
  story = []
  styles = getSampleStyleSheet()

  title_style = ParagraphStyle(
      "ChatTitle",
      parent=styles["Heading1"],
      fontSize=22,
      leading=26,
      textColor=HexColor("#1e293b"),
      spaceAfter=4,
  )
  sub_title_style = ParagraphStyle(
      "ChatSub",
      parent=styles["Normal"],
      fontSize=11,
      leading=15,
      textColor=HexColor("#64748b"),
      spaceAfter=15,
  )

  header_style_user = ParagraphStyle(
      "HeaderUser",
      parent=styles["Normal"],
      fontSize=11,
      leading=15,
      textColor=HexColor("#0284c7"),
      spaceBefore=6,
      spaceAfter=4,
  )
  header_style_tutor = ParagraphStyle(
      "HeaderTutor",
      parent=styles["Normal"],
      fontSize=11,
      leading=15,
      textColor=HexColor("#7e22ce"),
      spaceBefore=6,
      spaceAfter=4,
  )

  msg_style = ParagraphStyle(
      "MsgBody",
      parent=styles["Normal"],
      fontSize=10,
      leading=14,
      textColor=HexColor("#0f172a"),
      spaceAfter=6,
  )

  story.append(Paragraph("<b>STUDYSPHERE CHAT SESSION LOGS</b>", title_style))
  date_now = datetime.now().strftime("%Y-%m-%d")
  story.append(
      Paragraph(
          f"<b>Student:</b> {username} &nbsp;|&nbsp; <b>Context:</b>"
          f" {subject_title} &nbsp;|&nbsp; <b>Date:</b> {date_now}",
          sub_title_style,
      )
  )
  story.append(Spacer(1, 10))

  for msg in chat_history:
    if msg["role"] == "user":
      story.append(Paragraph(f"<b>🧑‍🎓 {username}</b>", header_style_user))
      bg_col = HexColor("#e0f2fe")
      border_col = HexColor("#0284c7")
    else:
      story.append(
          Paragraph("<b>⚡ StudySphere Tutor</b>", header_style_tutor)
      )
      bg_col = HexColor("#f8fafc")
      border_col = HexColor("#cbd5e1")

    paragraphs = [p.strip() for p in msg["content"].split("\n") if p.strip()]
    if not paragraphs:
      paragraphs = [msg["content"]]

    msg_flowables = []
    for p_text in paragraphs:
      clean_p = p_text.replace("<", "&lt;").replace(">", "&gt;")
      msg_flowables.append([Paragraph(clean_p, msg_style)])

    chunk_table = Table(msg_flowables, colWidths=[540])
    chunk_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), bg_col),
            ("BOX", (0, 0), (-1, -1), 1, border_col),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ])
    )

    story.append(chunk_table)
    story.append(Spacer(1, 8))

  doc.build(story)
  buffer.seek(0)
  return buffer.getvalue()


def display_report_card(username, subject_title, questions_scores):
  total_score = sum(q["score"] for q in questions_scores)
  max_score = sum(q["max_score"] for q in questions_scores)
  percentage = (total_score / max_score * 100) if max_score > 0 else 0
  grade = get_grade_letter(percentage)

  real_time_str = datetime.now().strftime("%Y-%m-%d")

  if "report_history" not in st.session_state:
    st.session_state.report_history = {}

  report_key = f"{subject_title} Report"
  st.session_state.report_history[report_key] = {
      "username": username,
      "subject": subject_title,
      "questions": questions_scores,
      "total": total_score,
      "max": max_score,
      "percentage": percentage,
      "grade": grade,
      "date": real_time_str,
  }

  pdf_bytes = generate_pdf_report_card(
      username,
      subject_title,
      questions_scores,
      total_score,
      max_score,
      percentage,
      grade,
      real_time_str,
  )

  st.download_button(
      label=f"📥 Download {subject_title} Report Card (PDF)",
      data=pdf_bytes,
      file_name=f"{subject_title.lower().replace(' ', '_')}_report_card.pdf",
      mime="application/pdf",
      use_container_width=True,
  )


# --- INITIALIZE STATE ---
saved_config = load_local_settings()

if "username" not in st.session_state:
  st.session_state.username = saved_config["user"]
if "questions_answered" not in st.session_state:
  st.session_state.questions_answered = saved_config["questions_answered"]
if "study_time_min" not in st.session_state:
  st.session_state.study_time_min = saved_config["study_time_min"]
if "uploaded_photo" not in st.session_state:
  if os.path.exists(PHOTO_FILE):
    try:
      with open(PHOTO_FILE, "rb") as f:
        st.session_state.uploaded_photo = f.read()
    except Exception:
      st.session_state.uploaded_photo = None
  else:
    st.session_state.uploaded_photo = None

if "chat_history" not in st.session_state:
  st.session_state.chat_history = []
if "assessment_questions" not in st.session_state:
  st.session_state.assessment_questions = []
if "assessment_evaluation" not in st.session_state:
  st.session_state.assessment_evaluation = None
if "mindmap_code" not in st.session_state:
  st.session_state.mindmap_code = None
if "flashcards_csv" not in st.session_state:
  st.session_state.flashcards_csv = None
if "report_history" not in st.session_state:
  st.session_state.report_history = {}

# --- SIDEBAR CONFIGURATION ---
with st.sidebar:
  if os.path.exists("logo.png"):
    _, center_col, _ = st.columns([1, 4, 1])
    with center_col:
      st.markdown('<div class="circle-logo-container">', unsafe_allow_html=True)
      st.image("logo.png", use_container_width=True)
      st.markdown("</div>", unsafe_allow_html=True)

  st.caption("// EXPAND YOUR MIND BEYOND BOUNDARIES")
  st.divider()

  with st.container(border=True):
    st.markdown("### 🧑‍🎓 Student Profile")
    col1, col2 = st.columns([1, 2])
    with col1:
      if st.session_state.uploaded_photo is not None:
        st.image(st.session_state.uploaded_photo, width=64)
      else:
        st.markdown("## ⚡")
    with col2:
      st.markdown(f"**{st.session_state.username}**")

  with st.expander("⚙️ Profile Configuration"):
    new_name = st.text_input("Username", value=st.session_state.username)
    photo_file = st.file_uploader(
        "Upload Profile Picture", type=["png", "jpg", "jpeg"]
    )
    if st.button("Save Profile Settings", use_container_width=True):
      st.session_state.username = new_name
      if photo_file is not None:
        photo_bytes = photo_file.read()
        st.session_state.uploaded_photo = photo_bytes
        try:
          with open(PHOTO_FILE, "wb") as f:
            f.write(photo_bytes)
        except Exception:
          pass
      save_local_settings(
          new_name,
          st.session_state.questions_answered,
          st.session_state.study_time_min,
      )
      st.success("Profile Updated!")
      st.rerun()

  st.divider()

  st.markdown("### 🎯 Learning Engine")
  study_mode = st.selectbox(
      "Select Active Mode:",
      [
          "📋 Practice Assessment & Evaluation",
          "📖 Lecture Mode (Deep Explanations)",
          "⚡ Flashcard Maker (Anki Export)",
          "🧠 Mind Map Node Visualizer",
          "👁️ Vision Notes & Formula Reader",
      ],
  )

  st.markdown("### 🎓 Academic Rigor Level")
  academic_level = st.segmented_control(
      "Academic Rigor Level",
      options=["High School", "Undergraduate", "Postgraduate / PhD"],
      default="Undergraduate",
      key="rigor_control",
      label_visibility="collapsed",
  )
  if not academic_level:
    academic_level = "Undergraduate"

  st.divider()

  st.markdown("### 📄 Multi-Format Context Material")
  uploaded_files = st.file_uploader(
      "Attach Course Notes (PDF, DOCX, PPTX, TXT)",
      type=["pdf", "docx", "pptx", "txt"],
      accept_multiple_files=True,
      key="doc_uploader",
  )

  document_text = ""
  active_subject_name = "Exam Assessment"
  if uploaded_files:
    active_subject_name = (
        os.path.splitext(uploaded_files[0].name)[0].replace("_", " ").title()
    )
    for f in uploaded_files:
      document_text += (
          f"\n--- File: {f.name} ---\n" + parse_uploaded_document(f)
      )
    if document_text:
      st.success(f"Ingested {len(uploaded_files)} Document(s)!")

  st.divider()

  st.markdown("### 📊 Cyber Analytics")
  st.caption(
      f"Questions Evaluated: **{st.session_state.questions_answered}**"
  )
  st.caption(
      "Est. Study Time Saved:"
      f" **{st.session_state.questions_answered * 4} mins**"
  )

  st.divider()

  if st.button("🧹 Clear Workspace", use_container_width=True):
    st.session_state.chat_history = []
    st.session_state.assessment_questions = []
    st.session_state.assessment_evaluation = None
    st.session_state.assessment_scores = None
    st.session_state.mindmap_code = None
    st.session_state.flashcards_csv = None
    st.success("Memory Purged!")
    st.rerun()

  with st.expander("💬 Session History Logs"):
    if st.session_state.chat_history:
      for msg in st.session_state.chat_history:
        sender = "User" if msg["role"] == "user" else "Tutor"
        st.caption(f"**{sender}:** {msg['content'][:80]}...")

      chat_pdf_data = generate_chat_pdf(
          username=st.session_state.username,
          subject_title=active_subject_name,
          chat_history=st.session_state.chat_history,
      )
      st.download_button(
          label="📥 Export Chat Session PDF",
          data=chat_pdf_data,
          file_name=(
              f"{active_subject_name.lower().replace(' ', '_')}_chat_logs.pdf"
          ),
          mime="application/pdf",
          use_container_width=True,
          key="sidebar_chat_pdf_download",
      )
    else:
      st.caption("No chat log entries yet.")

  with st.expander("📊 Grade Report Cards"):
    if st.session_state.report_history:
      selected_rep_sidebar = st.radio(
          "Saved Reports:",
          options=list(st.session_state.report_history.keys()),
          key="sidebar_report_selector",
      )
      if selected_rep_sidebar:
        rep_item = st.session_state.report_history[selected_rep_sidebar]
        st.caption(f"**Subject:** {rep_item['subject']}")
        st.caption(f"**Date:** {rep_item.get('date', 'N/A')}")
        st.caption(
            f"**Score:** {rep_item['total']}/{rep_item['max']}"
            f" ({rep_item['percentage']:.1f}%) | Grade: {rep_item['grade']}"
        )

        rep_pdf_bytes = generate_pdf_report_card(
            username=rep_item["username"],
            subject_title=rep_item["subject"],
            questions_scores=rep_item["questions"],
            total_score=rep_item["total"],
            max_score=rep_item["max"],
            percentage=rep_item["percentage"],
            grade=rep_item["grade"],
            real_time_str=rep_item.get(
                "date", datetime.now().strftime("%Y-%m-%d")
            ),
        )

        st.download_button(
            label=f"📥 Download PDF ({rep_item['subject']})",
            data=rep_pdf_bytes,
            file_name=(
                f"{rep_item['subject'].lower().replace(' ', '_')}_report_card.pdf"
            ),
            mime="application/pdf",
            use_container_width=True,
            key="sidebar_report_pdf_download",
        )
    else:
      st.caption("No saved report cards yet.")

# --- MAIN INTERFACE HEADER ---
st.markdown(
    "<h1 style='margin:0; padding-bottom: 4px;'>✨StudySphere</h1>",
    unsafe_allow_html=True,
)
st.caption(f"SYSTEM_MODE // {study_mode} | RIGOR // {academic_level.upper()}")
st.divider()

if not API_KEY:
  st.warning(
      "🔒 Security Notice: Gemini API Key is missing. Please configure"
      " `.streamlit/secrets.toml` or set the environment variable."
  )
else:
  client = None
  try:
    client = genai.Client(api_key=API_KEY)
  except Exception as e:
    st.error(f"Initialization Error: {e}")

  # ==========================================
  # MODE 1: PRACTICE ASSESSMENT & TIMED QUIZ
  # ==========================================
  if study_mode == "📋 Practice Assessment & Evaluation":
    st.subheader("📝 Interactive Assessment & Timed Quiz Mode")
    st.write("")

    col_type, col_timer = st.columns([1, 1], gap="medium")

    with col_type:
      with st.container(border=True):
        st.markdown("### 📑 Assessment Type")
        assessment_type = st.selectbox(
            "Choose Format:",
            [
                "Mixed Exam (MCQs & Written)",
                "Multiple Choice Questions (MCQ Only)",
                "Written & Short Answer Assessment",
                "Coding & Analytical Problem Solving",
            ],
            key="assessment_type_select",
        )

    with col_timer:
      with st.container(border=True):
        st.markdown("### ⏱️ Exam Duration")
        timer_options = {
            "10 Mins": 10 * 60,
            "30 Mins": 30 * 60,
            "1 Hr": 60 * 60,
            "3 Hrs": 180 * 60,
        }

        selected_timer_label = st.segmented_control(
            "Exam Duration",
            options=list(timer_options.keys()),
            default="30 Mins",
            key="timer_control",
            label_visibility="collapsed",
        )
        quiz_timer_sec = timer_options.get(selected_timer_label, 30 * 60)

    st.divider()

    if st.button("⚡ Generate Assessment Paper", use_container_width=True):
      if client:
        with st.spinner(f"Compiling {assessment_type} paper..."):
          try:
            if assessment_type == "Multiple Choice Questions (MCQ Only)":
              prompt_format = (
                  "Generate exactly 4 MCQ questions.\nOutput strictly JSON"
                  ' array of objects: [{"type": "mcq", "question": "...",'
                  ' "options": ["A)...", "B)...", "C)...", "D)..."]}]'
              )
            elif assessment_type == "Written & Short Answer Assessment":
              prompt_format = (
                  "Generate exactly 3 short-answer written conceptual"
                  ' questions.\nOutput strictly JSON array of objects: [{"type":'
                  ' "written", "question": "..."}]'
              )
            elif assessment_type == "Coding & Analytical Problem Solving":
              prompt_format = (
                  "Generate exactly 3 technical problem-solving or coding logic"
                  ' tasks.\nOutput strictly JSON array of objects: [{"type":'
                  ' "written", "question": "..."}]'
              )
            else:
              prompt_format = (
                  "Generate 2 MCQs and 2 short written questions.\nOutput"
                  " strictly JSON array of objects with keys 'type',"
                  " 'question', and optional 'options'."
              )

            prompt = (
                f"Generate an academic assessment paper for a {academic_level}"
                f" level.\nAssessment Style: {assessment_type}.\n{prompt_format}"
            )

            context_data = (
                f"DOCUMENT MATERIAL:\n{document_text}"
                if document_text
                else "General academic topics."
            )
            response = client.models.generate_content(
                model="gemini-2.5-flash",
                contents=[
                    f"Context: {context_data[:15000]}\n\nTask: {prompt}"
                ],
                config=types.GenerateContentConfig(
                    response_mime_type="application/json"
                ),
            )
            st.session_state.assessment_questions = json.loads(response.text)
            st.session_state.assessment_evaluation = None
            st.session_state.assessment_scores = None
            st.session_state.start_time = time.time()
            st.rerun()
          except Exception as e:
            st.error(handle_api_error(e))

    if st.session_state.assessment_questions:
      st.divider()

      if "start_time" in st.session_state:
        elapsed = time.time() - st.session_state.start_time
        remaining = max(0, int(quiz_timer_sec - elapsed))
        hours, remainder = divmod(remaining, 3600)
        mins, secs = divmod(remainder, 60)

        timer_str = (
            f"{hours:02d}:{mins:02d}:{secs:02d}"
            if hours > 0
            else f"{mins:02d}:{secs:02d}"
        )
        st.metric("⏳ Time Remaining", timer_str)
        if remaining == 0:
          st.warning(
              "⚠️ Time is up! Submit your answers now for evaluation."
          )

      st.subheader("✍️ Fill Answer Sheet")
      student_answers = []

      graded_map = {}
      if (
          "assessment_scores" in st.session_state
          and st.session_state.assessment_scores
      ):
        for score_item in st.session_state.assessment_scores:
          graded_map[score_item["q_num"]] = score_item

      for i, q_data in enumerate(st.session_state.assessment_questions):
        q_num = i + 1
        q_graded = graded_map.get(q_num)

        with st.container(border=True):
          if q_graded:
            is_corr = q_graded.get(
                "is_correct",
                q_graded.get("score", 0) >= (q_graded.get("max_score", 10) * 0.7),
            )
            badge_bg = (
                "rgba(34, 197, 94, 0.2)" if is_corr else "rgba(239, 68, 68, 0.2)"
            )
            badge_border = "#22c55e" if is_corr else "#ef4444"
            badge_text = "✅ CORRECT" if is_corr else "❌ WRONG"

            st.markdown(
                f"""
                            <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 8px;">
                                <h3 style="margin: 0; padding: 0; display: inline;">Question {q_num}</h3>
                                <span style="
                                    background: {badge_bg};
                                    border: 1px solid {badge_border};
                                    color: {badge_border};
                                    padding: 4px 12px;
                                    border-radius: 20px;
                                    font-weight: bold;
                                    font-size: 0.85rem;
                                ">{badge_text}</span>
                            </div>
                        """,
                unsafe_allow_html=True,
            )
          else:
            st.markdown(f"### Question {q_num}")

          st.markdown(f"**{q_data['question']}**")

          if q_data.get("type") == "mcq":
            selected_option = st.radio(
                "Select Option:", options=q_data["options"], key=f"mcq_{i}"
            )
            student_answers.append({
                "question": q_data["question"],
                "answer": selected_option,
                "type": "mcq",
            })
          else:
            ans_text = st.text_area(
                "Your Answer:",
                key=f"written_{i}",
                placeholder="Type your response or code solution here...",
            )
            student_answers.append({
                "question": q_data["question"],
                "answer": ans_text,
                "type": "written",
            })

          if q_graded and "explanation" in q_graded:
            st.caption(f"💡 **Feedback:** {q_graded['explanation']}")

      if st.button("🚀 Grade My Answers", use_container_width=True):
        if client:
          with st.spinner("Grading choices & written responses..."):
            try:
              eval_prompt = (
                  f"You are an examiner for a {academic_level} course. Evaluate"
                  " this exam sheet thoroughly.\n\n"
              )
              for idx, item in enumerate(student_answers):
                eval_prompt += (
                    f"Question {idx+1} ({item['type'].upper()}):"
                    f" {item['question']}\n"
                )
                eval_prompt += f"Student Answer {idx+1}: {item['answer']}\n\n"

              eval_prompt += (
                  "Provide structured scoring in JSON format:\nStrictly output"
                  " JSON array of objects with keys:\n- 'q_num' (integer 1, 2,"
                  " ...)\n- 'question_title' (short title string)\n- 'score'"
                  " (number 0-10)\n- 'max_score' (10)\n- 'is_correct' (boolean:"
                  " true if answer is correct, false if incorrect)\n-"
                  " 'explanation' (short 1-sentence feedback)\n\nExample:"
                  ' [{"q_num": 1, "question_title": "Concept A", "score": 10,'
                  ' "max_score": 10, "is_correct": true, "explanation":'
                  ' "Correct!"}]'
              )

              eval_response = client.models.generate_content(
                  model="gemini-2.5-flash",
                  contents=[eval_prompt],
                  config=types.GenerateContentConfig(
                      response_mime_type="application/json"
                  ),
              )

              st.session_state.assessment_scores = json.loads(eval_response.text)
              st.session_state.questions_answered += len(student_answers)
              save_local_settings(
                  st.session_state.username,
                  st.session_state.questions_answered,
                  st.session_state.study_time_min,
              )
              st.rerun()
            except Exception as e:
              st.error(handle_api_error(e))

    if (
        "assessment_scores" in st.session_state
        and st.session_state.assessment_scores
    ):
      st.divider()
      display_report_card(
          username=st.session_state.username,
          subject_title=active_subject_name,
          questions_scores=st.session_state.assessment_scores,
      )

  # ==========================================
  # MODE 2: FLASHCARD MAKER WITH ANKI CSV EXPORT
  # ==========================================
  elif study_mode == "⚡ Flashcard Maker (Anki Export)":
    st.subheader("⚡ Flashcard Deck Builder")
    topic_input = st.text_input(
        "Enter Topic or Concept for Flashcards:",
        placeholder="e.g., Quantum Mechanics Principles",
    )

    if st.button(
        "⚡ Generate Anki-Ready Flashcard Deck", use_container_width=True
    ):
      if client and (topic_input or document_text):
        with st.spinner("Generating flashcard pairs..."):
          try:
            prompt = (
                "Generate 5 high-yield flashcard Q&A pairs for a"
                f" {academic_level} student based on the context/topic.\nFormat"
                " output strictly as a JSON array of objects with keys: 'front'"
                " and 'back'."
            )
            ctx = (
                f"Topic: {topic_input}\nDocument Text:"
                f" {document_text[:10000]}"
            )
            response = client.models.generate_content(
                model="gemini-2.5-flash",
                contents=[f"{ctx}\n\n{prompt}"],
                config=types.GenerateContentConfig(
                    response_mime_type="application/json"
                ),
            )
            cards = json.loads(response.text)

            csv_lines = ["Front,Back"]
            for card in cards:
              f_txt = card["front"].replace('"', '""')
              b_txt = card["back"].replace('"', '""')
              csv_lines.append(f'"{f_txt}","{b_txt}"')

            st.session_state.flashcards_csv = "\n".join(csv_lines)
            st.session_state.flashcard_data = cards
          except Exception as e:
            st.error(handle_api_error(e))

    if "flashcard_data" in st.session_state and st.session_state.flashcard_data:
      st.divider()
      for idx, c in enumerate(st.session_state.flashcard_data):
        with st.expander(f"🎴 Flashcard {idx+1}: {c['front']}"):
          st.markdown(f"**Answer:** {c['back']}")

      st.download_button(
          label="📥 Export Flashcards for Anki / Quizlet (.CSV)",
          data=st.session_state.flashcards_csv,
          file_name="anki_flashcards.csv",
          mime="text/csv",
      )

  # ==========================================
  # MODE 3: MIND MAP NODE VISUALIZER (DIRECT GRAPHIC RENDER)
  # ==========================================
  elif study_mode == "🧠 Mind Map Node Visualizer":
    st.subheader("🧠 Interactive Concept Mind Map Generator")
    mindmap_topic = st.text_input(
        "Enter Concept to Map out:",
        placeholder="e.g., Neural Network Architecture",
    )

    if st.button("⚡ Generate Mind Map Diagram", use_container_width=True):
      if client and (mindmap_topic or document_text):
        with st.spinner("Mapping core nodes and compiling visual chart..."):
          try:
            prompt = (
                "Create a Graphviz DOT language graph representing the core"
                " concepts and sub-topics.\nRules:\n- Return ONLY valid DOT"
                " language syntax starting with 'digraph G {' and ending with"
                " '}'.\n- Do NOT use single quotes for attributes. Always use double quotes like bgcolor=\"#0a050d\" and color=\"#d946ef\".\n- Do not include markdown code block backticks or conversational text."
            )
            ctx = f"Topic: {mindmap_topic}\nMaterial: {document_text[:8000]}"
            response = client.models.generate_content(
                model="gemini-2.5-flash", contents=[f"{ctx}\n\n{prompt}"]
            )
            
            dot_code = clean_graphviz_dot(response.text)
            st.session_state.mindmap_code = dot_code

          except Exception as e:
            st.error(handle_api_error(e))

    if st.session_state.mindmap_code:
      st.divider()
      
      rendered = False
      
      # Try Primary Streamlit Native Renderer
      try:
        st.graphviz_chart(st.session_state.mindmap_code, use_container_width=True)
        rendered = True
      except Exception:
        pass

      # Try Backup Graphviz Source Pipe
      if not rendered:
        try:
          src = graphviz.Source(st.session_state.mindmap_code)
          st.image(src.pipe(format='png'), use_container_width=True)
          rendered = True
        except Exception:
          pass

      # Try Fallback Kroki SVG Web API
      if not rendered:
        try:
          url = "https://kroki.io/graphviz/svg/" + urllib.parse.quote(st.session_state.mindmap_code)
          st.image(url, use_container_width=True)
          rendered = True
        except Exception:
          st.error("Failed to render visual diagram. Please check DOT code below.")

      with st.expander("🔍 View Raw DOT Code (Developer View)"):
        st.code(st.session_state.mindmap_code, language="dot")

  # ==========================================
  # MODE 4: VISION NOTES & HANDWRITTEN READER
  # ==========================================
  elif study_mode == "👁️ Vision Notes & Formula Reader":
    st.subheader("👁️ Optical Note & Formula Breakdown")
    image_file = st.file_uploader(
        "Upload Image of Notes, Whiteboards, or Math Equations",
        type=["png", "jpg", "jpeg"],
    )

    if image_file and st.button(
        "⚡ Analyze & Explain Note Image", use_container_width=True
    ):
      if client:
        with st.spinner("Decoding image content with Gemini Vision..."):
          try:
            img_bytes = image_file.read()
            image_part = types.Part.from_bytes(
                data=img_bytes, mime_type=image_file.type
            )

            vision_prompt = (
                "Transcribe and thoroughly explain all handwritten notes,"
                f" formulas, or diagrams in this image for a {academic_level}"
                " student. Break down step-by-step solutions if math is present."
            )

            response = client.models.generate_content(
                model="gemini-2.5-flash", contents=[image_part, vision_prompt]
            )

            st.divider()
            st.markdown("### 🔍 Image Breakdown & Solution")
            st.markdown(response.text)

            tts = gTTS(text=response.text[:500], lang="en")
            speech_buffer = io.BytesIO()
            tts.write_to_fp(speech_buffer)
            st.audio(speech_buffer.getvalue(), format="audio/mp3")

          except Exception as e:
            st.error(handle_api_error(e))

  # ==========================================
  # MODE 5: LECTURE MODE & AUDIO TTS CHAT
  # ==========================================
  else:
    base_instruction = (
        "You are StudySphere, an expert engaging tutor. Tutor the student at a"
        f" {academic_level} level. Provide clear analogies and step-by-step"
        " points."
    )
    if document_text:
      base_instruction += (
          f"\n\nCONTEXT FROM UPLOADED DOCUMENTS:\n{document_text[:15000]}"
      )

    for message in st.session_state.chat_history:
      avatar_icon = "🧑‍🎓" if message["role"] == "user" else "⚡"
      with st.chat_message(message["role"], avatar=avatar_icon):
        st.markdown(message["content"])

    if user_prompt := st.chat_input(
        "Ask a question, enter a topic, or request an audio explanation..."
    ):
      st.session_state.chat_history.append(
          {"role": "user", "content": user_prompt}
      )

      with st.chat_message("user", avatar="🧑‍🎓"):
        st.markdown(user_prompt)

      with st.chat_message("model", avatar="⚡"):
        message_placeholder = st.empty()
        if not client:
          st.error("API Client missing or uninitialized.")
        else:
          try:
            formatted_contents = []
            for m in st.session_state.chat_history[-6:]:
              formatted_contents.append(
                  types.Content(
                      role=m["role"],
                      parts=[types.Part.from_text(text=m["content"])],
                  )
              )

            response = client.models.generate_content(
                model="gemini-2.5-flash",
                contents=formatted_contents,
                config=types.GenerateContentConfig(
                    system_instruction=base_instruction
                ),
            )

            tutor_response = response.text
            message_placeholder.markdown(tutor_response)
            st.session_state.chat_history.append(
                {"role": "model", "content": tutor_response}
            )

            tts = gTTS(text=tutor_response[:400], lang="en")
            audio_fp = io.BytesIO()
            tts.write_to_fp(audio_fp)
            st.audio(audio_fp.getvalue(), format="audio/mp3")

          except Exception as e:
            message_placeholder.error(handle_api_error(e))

    if st.session_state.chat_history:
      st.divider()

      chat_pdf_bytes = generate_chat_pdf(
          username=st.session_state.username,
          subject_title=active_subject_name,
          chat_history=st.session_state.chat_history,
      )

      st.download_button(
          label="📥 Export Chat Session (PDF)",
          data=chat_pdf_bytes,
          file_name=(
              f"{active_subject_name.lower().replace(' ', '_')}_chat_logs.pdf"
          ),
          mime="application/pdf",
          use_container_width=True,
      )
